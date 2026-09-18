#!/usr/bin/env python3
"""Build a source-linked, fail-closed experiment-team handoff.

All local-source and historical target/compound relationships remain in the full
matrix.  The workbook is an explicitly labeled working subset.  Literature
curation supplies reviewed claims, never a numerical efficacy or safety score.
"""

from __future__ import annotations

import argparse
import hashlib
import os
import json
import math
import re
import shutil
import sys
import zipfile
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

import pandas as pd
from rdkit import Chem
from rdkit.Chem.Scaffolds import MurckoScaffold

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from compound_applicability import assess
from reanalyze_target_candidate_safety import validated_model_artifact
from target_intent import (
    DECISION_SAFETY_KEYS,
    DECISION_SUMMARY_KEYS,
    build_target_intents,
    classify_evidence,
    decide_candidate,
)

SCHEMA = "skinscout.experiment-handoff.v1"
POLICY_SCOPES = {"intent_compound", "compound_global"}


def _private_input(env: str) -> Path | None:
    """비공개 입력은 환경변수로만 받는다(저장소에 경로를 남기지 않는다)."""
    value = os.environ.get(env, "").strip()
    return Path(value) if value else None
LABELS = {
    "experiment_priority": "실험 우선",
    "needs_evidence": "추가 근거 필요",
    "exclude": "목적 후보에서 제외",
    "unsupported": "현재 분석 지원 밖",
}


def digest(path: Path) -> str:
    result = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            result.update(block)
    return result.hexdigest()


def js(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def identity(smiles: str) -> dict:
    molecule = Chem.MolFromSmiles(smiles)
    if molecule is None:
        raise ValueError("Invalid exact compound SMILES")
    canonical = Chem.MolToSmiles(molecule, canonical=True, isomericSmiles=True)
    return {
        "compound_id": "structure-sha256:"
        + hashlib.sha256(canonical.encode("utf-8")).hexdigest(),
        "canonical_isomeric_smiles": canonical,
        "computed_inchikey": Chem.MolToInchiKey(molecule),
    }


def array(value: object) -> list:
    if isinstance(value, list):
        return value
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return []
    parsed = json.loads(str(value))
    if not isinstance(parsed, list):
        raise TypeError("Expected a JSON array")
    return parsed


def claim_digest(record: dict) -> str:
    return hashlib.sha256(
        json.dumps(
            record, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        ).encode("utf-8")
    ).hexdigest()


def load_curated(
    path: Path, literature: Path, intents: list[dict], review_path: Path | None = None
) -> list[dict]:
    registry = {row["intent_id"]: row for row in intents}
    metadata = {
        row["pmid"]: row
        for row in json.loads((literature / "sources/source_manifest.json").read_text())
    }
    payload = json.loads(path.read_text())
    if payload.get("schema_version") != "skinscout.curated-candidate-evidence.v1":
        raise ValueError("Unsupported literature curation schema")
    reviews = {}
    if review_path:
        review = json.loads(review_path.read_text())
        if review.get(
            "schema_version"
        ) != "skinscout.literature-review.v1" or review.get(
            "curation_sha256"
        ) != digest(path):
            raise ValueError("Independent literature review is missing or stale")
        for artifact in review["supporting_files"]:
            if digest(ROOT / artifact["path"]) != artifact["sha256"]:
                raise ValueError("Independent review supporting evidence changed")
        reviews = {row["evidence_id"]: row for row in review["records"]}
        if len(reviews) != len(review["records"]):
            raise ValueError("Duplicate independent-review evidence ID")
    rows = []
    seen = set()
    for record in payload["records"]:
        row = dict(record)
        if row["evidence_id"] in seen:
            raise ValueError("Duplicate curated evidence ID")
        seen.add(row["evidence_id"])
        intent = registry[row["intent_id"]]
        source = metadata[row["pmid"]]
        if source["doi"].lower() != row["doi"].lower():
            raise ValueError(f"DOI mismatch: {row['evidence_id']}")
        source_path = ROOT / source["path"]
        if digest(source_path) != source["sha256"]:
            raise ValueError(f"Changed literature source: {source_path}")
        identity_path = literature / "identity" / f"{row['identity_name']}.json"
        property_rows = json.loads(identity_path.read_text())["PropertyTable"][
            "Properties"
        ]
        if len(property_rows) != 1:
            raise ValueError("Ambiguous PubChem identity")
        properties = property_rows[0]
        if (
            properties["CID"] != row["expected_pubchem_cid"]
            or properties["InChIKey"] != row["expected_inchikey"]
        ):
            raise ValueError(f"PubChem identity mismatch: {row['compound_name']}")
        exact = identity(properties["SMILES"])
        if exact["computed_inchikey"] != row["expected_inchikey"]:
            raise ValueError("Isomeric graph and full InChIKey disagree")
        row.update(exact)
        reviewed = reviews.get(record["evidence_id"], {})
        if reviewed and (
            reviewed["record_sha256"] != claim_digest(record)
            or reviewed["source_sha256"] != source["sha256"]
        ):
            raise ValueError("Independent review no longer matches the claim/source")
        row["independent_review_verified"] = (
            reviewed.get("verdict") == "accepted_with_limitations"
        )
        row["target_intent_id"] = row["intent_id"]
        row.update(classify_evidence(intent, row))
        row.update(
            source_url=f"https://pubmed.ncbi.nlm.nih.gov/{row['pmid']}/",
            source_title=source["title"],
            source_sha256=source["sha256"],
            source_accessed_at=source["accessed_at"],
            identity_source_url=f"https://pubchem.ncbi.nlm.nih.gov/compound/{properties['CID']}",
            identity_source_sha256=digest(identity_path),
        )
        pmcid = source.get("pmcid")
        fulltext = literature / "sources" / f"{pmcid}.xml"
        if pmcid and fulltext.exists():
            row["fulltext_source_url"] = (
                f"https://pmc.ncbi.nlm.nih.gov/articles/{pmcid}/"
            )
            row["fulltext_sha256"] = digest(fulltext)
        rows.append(row)
    return rows


def validate_safety_model_evidence(row: dict, source_path: Path) -> None:
    """Reuse the safety producer's raw-artifact verification contract.

    The outer JSONL hash/count only proves the envelope; each recorded model
    response must still match its raw artifact bytes, schema, model, and
    request identity before the handoff consumes it.
    """
    model_results = row.get("model_results")
    if model_results is None:
        return
    if not isinstance(model_results, dict):
        raise ValueError(
            f"Safety record model_results must be an object: {source_path}"
        )
    for name, evidence in model_results.items():
        if not isinstance(evidence, dict):
            raise ValueError(
                f"Safety model evidence must be an object: {source_path}:{name}"
            )
        try:
            validated_model_artifact(name, evidence, row, source_path)
        except SystemExit as exc:
            raise ValueError(
                "Safety model evidence failed raw-artifact verification: "
                f"{source_path}:{name}: {exc}"
            ) from exc


def load_safety(paths: list[Path]) -> tuple[dict[str, dict], list[dict]]:
    records: dict[str, dict] = {}
    artifacts = []
    for path in paths:
        manifest_path = path.parent / "manifest.json"
        manifest = json.loads(manifest_path.read_text())
        expected = manifest["artifacts"]["candidate_safety_jsonl"]["sha256"]
        if digest(path) != expected:
            raise ValueError(f"Safety input hash mismatch: {path}")
        rows = [
            json.loads(line) for line in path.read_text().splitlines() if line.strip()
        ]
        expected_count = manifest.get(
            "candidate_count", manifest.get("selected_unique_compound_count")
        )
        if len(rows) != expected_count:
            raise ValueError("Safety manifest count mismatch")
        for row in rows:
            exact = identity(row["canonical_smiles"])
            if exact["compound_id"] != row["compound_id"]:
                raise ValueError("Safety record belongs to another exact structure")
            validate_safety_model_evidence(row, path)
            previous = records.get(row["compound_id"])
            if previous and previous != row:
                raise ValueError(
                    "Nonidentical duplicate safety records require explicit reconciliation"
                )
            records[row["compound_id"]] = row
        artifacts.append(
            {"path": str(path), "sha256": digest(path), "manifest": manifest}
        )
    return records, artifacts


def aggregate_direction(values: list[str]) -> tuple[str, str]:
    meaningful = set(values) & {"supports", "opposes", "conflict"}
    if "conflict" in meaningful:
        return "conflict", "A reviewed comparable-context conflict is present."
    if {"supports", "opposes"}.issubset(meaningful):
        return (
            "unknown",
            "Mixed directions across records; assay-context adjudication is required.",
        )
    if meaningful:
        return (
            next(iter(meaningful)),
            "Recorded direction; human relevance and evidence completeness are separate gates.",
        )
    return "unknown", "No eligible directional record."


def build_matrix(
    intents: list[dict],
    source_summary: pd.DataFrame,
    historical: pd.DataFrame,
    curated: list[dict],
    safety: dict[str, dict],
    reference_controls: list[dict] | None = None,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    intent_map = {row["intent_id"]: row for row in intents}
    uniprot_map: dict[str, list[str]] = {}
    for row in intents:
        if row["uniprot_id"]:
            uniprot_map.setdefault(row["uniprot_id"], []).append(row["intent_id"])
    pairs: dict[tuple[str, str], dict] = {}
    compounds: dict[str, dict] = {}
    controls = {
        (row["intent_id"], row["compound_id"]): row
        for row in (reference_controls or [])
    }
    global_policy_by_compound: dict[str, list[dict]] = {}
    for row in curated:
        scope = row.get("policy_scope")
        if scope is not None and str(scope).strip() not in POLICY_SCOPES:
            raise ValueError(f"Invalid material policy scope: {scope!r}")
        if str(scope or "").strip() == "compound_global":
            global_policy_by_compound.setdefault(row["compound_id"], []).append(row)

    def put(intent_id: str, exact: dict, name: str | None = None) -> dict:
        key = (intent_id, exact["compound_id"])
        if intent_id not in intent_map:
            raise ValueError(f"Unmapped intent {intent_id}")
        existing = compounds.get(exact["compound_id"])
        if (
            existing
            and existing["canonical_isomeric_smiles"]
            != exact["canonical_isomeric_smiles"]
        ):
            raise ValueError("Compound ID collision")
        if not existing:
            compounds[exact["compound_id"]] = {
                **exact,
                "compound_name": name
                or exact["computed_inchikey"]
                or exact["compound_id"],
            }
        elif name:
            existing["compound_name"] = name
        return pairs.setdefault(
            key,
            {
                "intent_id": intent_id,
                "compound_id": exact["compound_id"],
                "source_evidence_count": 0,
                "source_publication_count": 0,
                "source_direction_values": [],
                "source_identity_status": None,
                "source_identity_statuses": [],
                "source_identity_mismatch_count": 0,
                "historical_rows": [],
                "curated": [],
            },
        )

    for row in source_summary.to_dict("records"):
        if not row.get("canonical_isomeric_smiles") or pd.isna(
            row["canonical_isomeric_smiles"]
        ):
            if not re.fullmatch(
                r"unresolved-(?:structure|source-compound)-sha256:[0-9a-f]{64}",
                str(row["compound_id"]),
            ):
                raise ValueError(
                    "Missing structure has no explicit unresolved identity"
                )
            exact = {
                "compound_id": row["compound_id"],
                "canonical_isomeric_smiles": None,
                "computed_inchikey": None,
            }
        else:
            exact = identity(row["canonical_isomeric_smiles"])
            if exact["compound_id"] != row["compound_id"]:
                raise ValueError("Source summary compound ID mismatch")
        name = row.get("molecule_pref_name")
        if pd.isna(name):
            name = None
        pair = put(row["intent_id"], exact, name)
        if pair["source_evidence_count"]:
            raise ValueError("Duplicate source summary intent/compound pair")
        pair["source_evidence_count"] = int(row["evidence_count"])
        pair["source_publication_count"] = int(row["unique_publication_count"])
        pair["source_direction_values"] = array(row["direction_relations_json"])
        pair["source_identity_status"] = row.get("structure_identity_status")
        pair["source_identity_statuses"] = array(
            row.get("structure_identity_statuses_json")
        )
        pair["source_identity_mismatch_count"] = int(
            row.get("source_identity_mismatch_count", 0)
        )
    for number, row in enumerate(historical.to_dict("records"), 2):
        matching = uniprot_map.get(str(row["uniprot"]), [])
        if not matching:
            raise ValueError(f"Historical target is unmapped: {row['uniprot']}")
        exact = identity(row["smiles"])
        for intent_id in matching:
            pair = put(intent_id, exact)
            pair["historical_rows"].append(number)
    for row in curated:
        exact = {
            key: row[key]
            for key in ("compound_id", "canonical_isomeric_smiles", "computed_inchikey")
        }
        pair = put(row["intent_id"], exact, row["compound_name"])
        pair["curated"].append(row)
    for control in controls.values():
        exact = {
            key: control[key]
            for key in ("compound_id", "canonical_isomeric_smiles", "computed_inchikey")
        }
        put(control["intent_id"], exact, control["source_compound_id"])

    for compound_id, row in compounds.items():
        scope = assess(row["canonical_isomeric_smiles"] or "")
        current = safety.get(compound_id, {})
        if current and current["applicability"]["verdict"] != scope["verdict"]:
            raise ValueError(
                "Current applicability code disagrees with the safety artifact"
            )
        row.update(
            applicability=scope["verdict"],
            applicability_exclusions=js(scope["exclusions"]),
            applicability_warnings=js(scope["warnings"]),
            molecular_weight=scope.get("properties", {}).get("molecular_weight"),
            safety_record=bool(current),
            web_model_attempted=bool(current.get("model_results")),
            safety_decision=current.get("decision", "NOT_RUN"),
            safety_run_valid=current.get("run_valid", False),
            safety_consensus_valid=current.get("safety_consensus_valid", False),
            full_analysis_complete=current.get("full_analysis_complete", False),
            safety_domain=current.get("applicability_domain", "unknown"),
            safety_degraded=current.get("skin_sens", {}).get("degraded", True),
            safety_calls=js(current.get("skin_sens", {}).get("calls", {})),
            safety_missing_models=js(
                current.get("skin_sens", {}).get("missing_models", [])
            ),
            admet_status=current.get("admet_status", "not_run"),
            old_safety_decision=current.get("historical", {}).get("old_skin_sens"),
            model_execution=js(
                {
                    name: result.get("execution", "unknown")
                    for name, result in current.get("model_results", {}).items()
                }
            ),
        )
    rows = []
    for key in sorted(pairs):
        pair = pairs[key]
        intent = intent_map[pair["intent_id"]]
        compound = compounds[pair["compound_id"]]
        papers = pair["curated"]
        values = pair["source_direction_values"] + [
            row["direction_relation"] for row in papers
        ]
        direction, context = aggregate_direction(values)
        qualified = [
            row
            for row in papers
            if row["primary_evidence_verified"] is True
            and row["human_relevance_verified"] is True
            and row.get("independent_review_verified") is True
            and row["direction_relation"] == "supports"
        ]
        # Role and material policy are scoped to this intent+compound pair.
        # A compound-global policy only applies when a curated record declares
        # that scope with an explicit basis, so curating one intent cannot
        # silently change another intent of the same compound.
        global_papers = global_policy_by_compound.get(pair["compound_id"], [])
        if global_papers:
            policy_papers = global_papers
            policy_scope = "compound_global"
            bases = sorted({
                str(row.get("policy_basis", "")).strip() for row in global_papers
            })
            if len(bases) != 1 or not bases[0]:
                raise ValueError(
                    "Compound-global role/material policy requires one explicit basis"
                )
            policy_basis = bases[0]
        else:
            policy_papers = papers
            policy_scope = "intent_compound"
            policy_basis = "; ".join(
                sorted(f"curated:{row['evidence_id']}" for row in policy_papers)
            ) or "no curated evidence for this intent-compound pair"
        role = "material_candidate"
        if policy_papers and all(
            row["role"] == "reference_control" for row in policy_papers
        ):
            role = "reference_control"
        elif policy_papers and all(
            row["role"] == "exploratory_candidate" for row in policy_papers
        ):
            role = "exploratory_candidate"
        policy_values = {row["material_policy"] for row in policy_papers}
        material_policy = (
            "excluded"
            if "excluded" in policy_values
            else ("eligible" if policy_values == {"eligible"} else "unknown")
        )
        control = controls.get((pair["intent_id"], pair["compound_id"]))
        if control is not None:
            role = "reference_control"
            material_policy = "excluded"
            policy_scope = "intent_compound"
            policy_basis = (
                f"structural control pair {control.get('pair_id', '')}".strip()
            )
        # The raw applicability verdict stays a separate input from the computed
        # support flag: `review` used to be rewritten to `in_scope` here, so a
        # panel-range caution reached the same decision path as an in-scope
        # compound while the raw state survived only in the display columns.
        # "supported" only means the analysis is inside the documented scope
        # (not an exclusion); `review` is passed through raw and still forces
        # needs_evidence in decide_candidate.
        raw_applicability = str(compound["applicability"])
        analysis_scope_supported = raw_applicability in {"in_scope", "review"}
        applicability_warning_codes = [
            str(item.get("code"))
            for item in array(compound.get("applicability_warnings"))
            if isinstance(item, dict)
        ]
        summary = {
            "structure_valid": bool(compound["canonical_isomeric_smiles"]),
            "route_supported": True,
            "analysis_applicability": raw_applicability,
            "analysis_scope_supported": analysis_scope_supported,
            "direction_relation": direction,
            "material_policy": material_policy,
            "required_information_complete": any(
                row["conditions_complete"] for row in qualified
            ),
            "minimum_evidence_met": bool(qualified),
            "functional_evidence": bool(qualified),
        }
        current = safety.get(compound["compound_id"], {})
        verdict = decide_candidate(intent, summary, current)
        decision_inputs = {
            "applicability": {
                "raw_state": raw_applicability,
                "warning_codes": applicability_warning_codes,
                "analysis_scope_supported": analysis_scope_supported,
            },
            "summary": {key: summary.get(key) for key in DECISION_SUMMARY_KEYS},
            "safety": {key: current.get(key) for key in DECISION_SAFETY_KEYS},
        }
        structural = (
            "not_applicable_endpoint"
            if intent["route"] == "endpoint"
            else "not_applicable_pathway"
            if not intent["docking_eligible"]
            else "abstained_unvalidated_metal_support"
            if intent["gene_symbol"] in {"TYR", "MMP1", "MMP3"}
            else "not_run"
        )
        rows.append(
            {
                **{
                    k: intent[k]
                    for k in (
                        "intent_id",
                        "source_row_id",
                        "category",
                        "biomarker",
                        "gene_symbol",
                        "uniprot_id",
                        "route",
                        "desired_effect",
                        "skin_compartment",
                        "readout",
                    )
                },
                **compound,
                "role": role,
                "decision": verdict["decision"],
                "decision_label": LABELS[verdict["decision"]],
                "decision_reasons": js(verdict["reasons"]),
                "analysis_applicability": raw_applicability,
                "analysis_scope_supported": analysis_scope_supported,
                "decision_inputs": js(decision_inputs),
                "direction_relation": direction,
                "direction_context": context,
                "minimum_evidence_met": summary["minimum_evidence_met"],
                "conditions_complete": summary["required_information_complete"],
                "material_policy": material_policy,
                "material_policy_scope": policy_scope,
                "material_policy_basis": policy_basis,
                "role_scope": policy_scope,
                "source_evidence_count": pair["source_evidence_count"],
                "source_identity_status": pair["source_identity_status"],
                "source_identity_statuses": js(pair["source_identity_statuses"]),
                "source_identity_mismatch_count": pair[
                    "source_identity_mismatch_count"
                ],
                "source_publication_count_unadjudicated": pair[
                    "source_publication_count"
                ],
                "reviewed_primary_publication_count": len(
                    {r["pmid"] for r in qualified}
                ),
                "curated_evidence_ids": js([r["evidence_id"] for r in papers]),
                "source_urls": js(sorted({r["source_url"] for r in papers})),
                "reviewed_evidence": " | ".join(r["reported_result"] for r in papers),
                "reviewed_conditions": " | ".join(
                    r["reported_conditions"] or "unknown" for r in papers
                ),
                "evidence_limitations": " | ".join(r["limitations"] for r in papers),
                "next_experiment": " | ".join(
                    dict.fromkeys(r["next_experiment"] for r in papers)
                )
                or "Review original functional evidence and qualify the assay before candidate testing.",
                "historical_relationship_rows": js(pair["historical_rows"]),
                "is_curated_working_set": bool(papers),
                "structure_status": structural,
                "procurement_status": "not_purchased; supplier, lot, purity and exact isomer unverified",
                "formulation_exposure_status": "unassessed",
            }
        )
    return pd.DataFrame(rows), pd.DataFrame(list(compounds.values()))


def historical_counts(
    intents: list[dict], historical: pd.DataFrame, matrix: pd.DataFrame
) -> dict:
    graphs = {
        index + 2: identity(row["smiles"])["compound_id"]
        for index, row in enumerate(historical.to_dict("records"))
    }
    expected = {}
    for index, row in enumerate(historical.to_dict("records"), 2):
        expected[index] = {
            (i["intent_id"], graphs[index])
            for i in intents
            if i["uniprot_id"] == row["uniprot"]
        }
    observed = {index: set() for index in expected}
    mentions = Counter()
    for row in matrix.itertuples():
        for number in array(row.historical_relationship_rows):
            if number not in observed:
                raise ValueError("Unknown historical source row in matrix")
            observed[number].add((row.intent_id, row.compound_id))
            mentions[number] += 1
    if expected != observed or any(
        mentions[key] != len(value) for key, value in expected.items()
    ):
        raise ValueError("Historical target/compound lineage was lost or duplicated")
    return {
        "historical_raw_rows": len(historical),
        "historical_unique_exact_compounds": len(set(graphs.values())),
        "historical_unique_intent_compound_pairs": len(set().union(*expected.values()))
        if expected
        else 0,
        "historical_lineage_mentions": sum(mentions.values()),
    }


STRUCTURAL_SCHEMA = "skinscout.target-candidate-structural-screen.v1"
STRUCTURAL_SCORE_FIELDS = (
    "gnina_cnn_affinity",
    "gnina_cnn_score",
    "gnina_minimized_affinity",
)
STRUCTURAL_BYTE_FIELDS = (
    ("pose_sdf", "pose_sha256"),
    ("ligand_sdf", "ligand_sha256"),
    ("receptor_pdb", "receptor_sha256"),
    ("pocket_json", "pocket_sha256"),
)
STRUCTURAL_SHARED_FIELDS = (
    "intent_id",
    "target_id",
    "compound_id",
    "control_role",
    "control_evidence_id",
    "status",
    "safety_status",
    "applicability_status",
    "safety_consensus_valid",
    "full_analysis_complete",
    "structure_protocol_status",
    "structure_protocol_evidence_sha256",
    "safety_evidence_sha256",
    "safety_evidence_record_sha256",
    "ligand_graph_smiles",
    "ligand_inchi_key",
    "pose_sha256",
) + STRUCTURAL_SCORE_FIELDS


def canonical_json_sha256(value: object) -> str:
    return hashlib.sha256(
        json.dumps(
            value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        ).encode("utf-8")
    ).hexdigest()


def manifest_field_absent(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def structural_input_path(directory: Path, raw: object) -> Path:
    """Resolve a manifest path, tolerating portable copies that renamed files."""
    candidate = Path(str(raw))
    for path in (candidate, directory / candidate.name):
        if path.is_file():
            return path
    matches = sorted(
        path for path in directory.rglob(candidate.name) if path.is_file()
    ) if candidate.name else []
    if len(matches) == 1:
        return matches[0]
    return candidate


def same_number(left: object, right: object) -> bool:
    if (left is None or left == "") and (right is None or right == ""):
        return True
    try:
        left_value, right_value = float(left), float(right)
    except (TypeError, ValueError):
        return False
    return (
        math.isfinite(left_value)
        and math.isfinite(right_value)
        and left_value == right_value
    )


def validate_structural_manifest(directory: Path) -> tuple[dict, dict[str, dict]]:
    """Bind pair_status.csv to the sealed structural screen manifest.

    Every pair must join 1:1 by pair_id, agree on identity/status/scores/hashes,
    and have pose and source bytes matching the recorded SHA-256.  This rejects
    a later copy that only resembles the original run.
    """
    manifest_path = directory / "structural_manifest.json"
    if not manifest_path.is_file():
        raise ValueError(f"Structural screen manifest is required: {manifest_path}")
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"Structural screen manifest failed to parse: {manifest_path}: {exc}"
        ) from exc
    if (
        not isinstance(manifest, dict)
        or manifest.get("schema_version") != STRUCTURAL_SCHEMA
    ):
        raise ValueError("Unsupported structural screen manifest schema")
    for name in ("pair_source", "source_record_manifest"):
        reference = manifest.get(name)
        if not isinstance(reference, dict) or not isinstance(
            reference.get("sha256"), str
        ):
            raise ValueError(f"Structural manifest has no {name} binding")  # noqa: TRY004
        path = structural_input_path(directory, reference.get("path"))
        if not path.is_file():
            raise ValueError(f"Structural {name} input is missing: {path}")
        if digest(path) != reference["sha256"]:
            raise ValueError(f"Structural {name} hash mismatch: {path}")
    pairs = manifest.get("pairs")
    if not isinstance(pairs, list) or not pairs:
        raise ValueError("Structural manifest has no pairs")
    manifest_by_id: dict[str, dict] = {}
    for pair in pairs:
        if not isinstance(pair, dict):
            raise ValueError("Structural manifest pair is not an object")  # noqa: TRY004
        pair_id = str(pair.get("pair_id", "")).strip()
        if not pair_id or pair_id in manifest_by_id:
            raise ValueError(f"Structural manifest pair_id is blank or duplicated: {pair_id!r}")
        manifest_by_id[pair_id] = pair
    status_path = directory / "pair_status.csv"
    if not status_path.is_file():
        raise ValueError(f"Structural pair status CSV is required: {status_path}")
    status_by_id: dict[str, dict] = {}
    for row in pd.read_csv(
        status_path, dtype=str, keep_default_na=False
    ).to_dict("records"):
        pair_id = str(row.get("pair_id", "")).strip()
        if not pair_id or pair_id in status_by_id:
            raise ValueError(f"pair_status.csv pair_id is blank or duplicated: {pair_id!r}")
        status_by_id[pair_id] = row
    if set(manifest_by_id) != set(status_by_id):
        raise ValueError(
            "Structural manifest and pair_status.csv must join 1:1 by pair_id"
        )
    for pair_id, pair in manifest_by_id.items():
        status = status_by_id[pair_id]
        for field in STRUCTURAL_SHARED_FIELDS:
            if field not in pair or field not in status:
                continue
            if field in STRUCTURAL_SCORE_FIELDS:
                if not same_number(pair[field], status[field]):
                    raise ValueError(
                        f"Structural pair {pair_id} {field} disagrees with pair_status.csv"
                    )
            elif manifest_field_absent(pair[field]) and manifest_field_absent(
                status[field]
            ):
                continue
            elif str(pair[field]) != str(status[field]):
                raise ValueError(
                    f"Structural pair {pair_id} {field} disagrees with pair_status.csv"
                )
        status_text = str(pair.get("status", "")).strip()
        if not status_text:
            raise ValueError(f"Structural pair {pair_id} has no status")
        if status_text == "completed":
            pose_path = structural_input_path(directory, pair.get("pose_sdf"))
            if not pose_path.is_file() or digest(pose_path) != pair.get("pose_sha256"):
                raise ValueError(f"Structural pair {pair_id} pose hash mismatch")
            for path_field, hash_field in STRUCTURAL_BYTE_FIELDS[1:]:
                path = structural_input_path(directory, pair.get(path_field))
                if not path.is_file() or digest(path) != pair.get(hash_field):
                    raise ValueError(
                        f"Structural pair {pair_id} {path_field} hash mismatch"
                    )
        else:
            reason = str(pair.get("reason") or "").strip()
            if not reason:
                raise ValueError(
                    f"Structural pair {pair_id} status {status_text!r} "
                    "requires a reason"
                )
            recorded = [
                field
                for path_field, hash_field in STRUCTURAL_BYTE_FIELDS
                for field in (path_field, hash_field)
                if not manifest_field_absent(pair.get(field))
            ]
            if recorded:
                raise ValueError(
                    f"Structural pair {pair_id} status {status_text!r} must not "
                    f"record structural bytes: {', '.join(recorded)}"
                )
    return manifest, status_by_id


def load_structural_controls(
    directory: Path | None,
) -> tuple[list[dict], pd.DataFrame, dict]:
    if directory is None:
        return [], pd.DataFrame(), {"status": "not_provided", "pairs": []}
    manifest, _ = validate_structural_manifest(directory)
    source_path = structural_input_path(
        directory, manifest["source_record_manifest"]["path"]
    )
    source_manifest = json.loads(source_path.read_text(encoding="utf-8"))
    source_controls = source_manifest.get("controls")
    if not isinstance(source_controls, list):
        raise ValueError(  # noqa: TRY004
            "Structural source-record manifest has no controls list"
        )
    by_compound: dict[str, dict] = {}
    for control in source_controls:
        if not isinstance(control, dict):
            raise ValueError("Structural source control is not an object")  # noqa: TRY004
        compound = control.get("structure_compound_id")
        if not isinstance(compound, str) or not compound or compound in by_compound:
            raise ValueError("Structural source controls require unique compound ids")
        by_compound[compound] = control
    controls = []
    for pair in manifest["pairs"]:
        pair_id = pair["pair_id"]
        source = by_compound.get(pair["compound_id"])
        if source is None:
            raise ValueError(f"Structural pair {pair_id} has no source control record")
        if canonical_json_sha256(source) != pair.get("source_control_record_sha256"):
            raise ValueError(f"Structural pair {pair_id} source control hash mismatch")
        exact = identity(source["smiles"])
        if exact["computed_inchikey"] != source["standard_inchi_key"]:
            raise ValueError("Structural control stereochemical identity mismatch")
        if str(pair.get("status", "")).strip() == "completed":
            ligand = structural_input_path(directory, source["ligand_sdf"])
            if digest(ligand) != source["ligand_sdf_sha256"]:
                raise ValueError("Structural control ligand file was changed")
        controls.append(
            {
                **exact,
                "pair_id": pair_id,
                "intent_id": pair["intent_id"],
                "source_compound_id": source["molecule_chembl_id"],
                "control_role": pair["control_role"],
                "assay_id": source["assay_chembl_id"],
                "activity_id": source["activity_id"],
                "reported_type": source["standard_type"],
                "reported_relation": source["standard_relation"],
                "reported_value": source["standard_value"],
                "reported_units": source["standard_units"],
                "source_url": "https://doi.org/" + source["doi"],
                "structure_status": pair["status"],
                "structure_reason": pair.get("reason"),
                "safety_status": pair.get("safety_status"),
                "gnina_cnn_affinity": pair.get("gnina_cnn_affinity"),
                "gnina_cnn_score": pair.get("gnina_cnn_score"),
                "gnina_minimized_affinity": pair.get("gnina_minimized_affinity"),
                "interpretation": "Reference control only. Weak binding is not inactivity. Structural scores do not establish functional direction or safety.",
            }
        )
    artifacts = {
        "status": "verified",
        "manifest": {
            "path": str(directory / "structural_manifest.json"),
            "sha256": digest(directory / "structural_manifest.json"),
        },
        "pair_status": {
            "path": str(directory / "pair_status.csv"),
            "sha256": digest(directory / "pair_status.csv"),
        },
        "source_record_manifest": {
            "path": str(source_path),
            "sha256": digest(source_path),
        },
        "pair_source": {
            "path": str(
                structural_input_path(directory, manifest["pair_source"]["path"])
            ),
            "sha256": manifest["pair_source"]["sha256"],
        },
        "pairs": [
            {
                "pair_id": pair["pair_id"],
                "status": pair["status"],
                "target_id": pair.get("target_id"),
                "compound_id": pair.get("compound_id"),
                "pose_sha256": pair.get("pose_sha256"),
                "ligand_sha256": pair.get("ligand_sha256"),
                "receptor_sha256": pair.get("receptor_sha256"),
                "pocket_sha256": pair.get("pocket_sha256"),
                "safety_evidence_sha256": pair.get("safety_evidence_sha256"),
                "source_control_record_sha256": pair.get("source_control_record_sha256"),
            }
            for pair in manifest["pairs"]
        ],
    }
    return controls, pd.DataFrame(controls), artifacts


def select_priority(matrix: pd.DataFrame, top: int = 3) -> pd.DataFrame:
    eligible = matrix[matrix.decision.eq("experiment_priority")].copy()
    eligible = eligible.sort_values(
        ["intent_id", "reviewed_primary_publication_count", "compound_id"],
        ascending=[True, False, True],
    )
    indices = []
    for _, group in eligible.groupby("intent_id", sort=True):
        seen = set()
        for index, row in group.iterrows():
            molecule = Chem.MolFromSmiles(row.canonical_isomeric_smiles)
            scaffold = MurckoScaffold.MurckoScaffoldSmiles(
                mol=molecule, includeChirality=True
            )
            scaffold_key = scaffold or row.compound_id
            if scaffold_key not in seen:
                seen.add(scaffold_key)
                indices.append(index)
            if len(seen) >= top:
                break
    return eligible.loc[indices].copy()


def portable_relative(path: Path) -> Path:
    if path.is_absolute() or ".." in path.parts:
        raise ValueError("Artifact path must remain relative to the package")
    reserved = {"CON", "PRN", "AUX", "NUL"} | {
        f"{prefix}{number}" for prefix in ("COM", "LPT") for number in range(1, 10)
    }
    parts = []
    for part in path.parts:
        clean = re.sub(r'[<>:"\\|?*\x00-\x1f]', "_", part).rstrip(" .") or "_"
        if clean.split(".")[0].upper() in reserved:
            clean = "_" + clean
        parts.append(clean)
    return Path(*parts)


def copy_artifact_tree(
    source: Path, destination: Path, package_root: Path, path_map: list[dict]
) -> None:
    copied = set()
    for path in sorted(source.rglob("*")):
        if not path.is_file():
            continue
        relative = portable_relative(path.relative_to(source))
        collision_key = str(relative).casefold()
        if collision_key in copied:
            raise ValueError("Portable artifact paths collide")
        copied.add(collision_key)
        target = destination / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, target)
        resolved = path.resolve()
        path_map.append(
            {
                "workspace_path": str(resolved.relative_to(ROOT))
                if resolved.is_relative_to(ROOT)
                else str(resolved),
                "package_path": str(target.relative_to(package_root)),
                "sha256": digest(target),
            }
        )


def control_ranking_summary(controls: pd.DataFrame) -> dict:
    if controls.empty or not controls.structure_status.eq("completed").all():
        return {"status": "not_evaluated"}
    active = controls[controls.control_role.eq("active_control")]
    weak = controls[controls.control_role.eq("weak_binding_control")]
    if len(active) != 1 or len(weak) != 1:
        return {"status": "not_evaluated"}
    first, second = active.iloc[0], weak.iloc[0]
    matches = {
        "gnina_cnn_affinity": None,
        "gnina_cnn_score": None,
        "gnina_minimized_affinity": None,
    }
    if not (
        first.assay_id == second.assay_id
        and first.reported_type == second.reported_type == "Ki"
        and first.reported_units == second.reported_units
        and first.reported_relation == second.reported_relation == "="
        and first.reported_value < second.reported_value
    ):
        return {"status": "not_comparable"}
    for field in matches:
        if (
            pd.notna(first[field])
            and pd.notna(second[field])
            and first[field] != second[field]
        ):
            matches[field] = (
                bool(first[field] < second[field])
                if field == "gnina_minimized_affinity"
                else bool(first[field] > second[field])
            )
    return {
        "status": "single_control_pair_only",
        "metric_order_matches_measured_affinity": matches,
        "interpretation": "One control pair is not a ranking benchmark; no score is selected after observing these results.",
    }


SAFETY_COVERAGE_COLUMNS = {
    "safety_record_count": "safety_record",
    "web_model_attempted_count": "web_model_attempted",
    "safety_consensus_valid_count": "safety_consensus_valid",
    "full_analysis_complete_count": "full_analysis_complete",
}


def safety_coverage_metrics(matrix: pd.DataFrame) -> dict[str, dict[str, int]]:
    """Count safety states separately at pair and exact-compound units.

    A stored record, an attempted web-model run, a valid consensus (which can
    include a partial two-positive HALT), and a completed analysis are distinct
    states.  Each metric keeps the intent-compound pair and unique exact
    compound denominators so they are never compared as the same unit.
    """
    metrics: dict[str, dict[str, int]] = {}
    for name, column in SAFETY_COVERAGE_COLUMNS.items():
        present = matrix[column].fillna(False).astype(bool)
        metrics[name] = {
            "unit": "intent-compound pairs / exact compounds",
            "pairs": int(present.sum()),
            "compounds": int(matrix.loc[present, "compound_id"].nunique()),
        }
    return metrics


def coverage_table(intents: list[dict], matrix: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for intent in intents:
        group = matrix[matrix.intent_id.eq(intent["intent_id"])]
        states = Counter(group.decision)
        rows.append(
            {
                **{
                    key: intent[key]
                    for key in (
                        "source_row_id",
                        "source_row_number",
                        "intent_id",
                        "category",
                        "biomarker",
                        "gene_symbol",
                        "uniprot_id",
                        "route",
                        "desired_effect",
                        "readout",
                    )
                },
                "source_lookup_completed": intent["route"] != "endpoint",
                "source_lookup_status": "not_applicable_nonprotein_endpoint"
                if intent["route"] == "endpoint"
                else "complete_local_source_lookup",
                "source_evidence_rows": int(group.source_evidence_count.sum()),
                "candidate_count": len(group),
                "curated_pair_count": int(group.is_curated_working_set.sum()),
                "safety_record_count": int(group.safety_record.astype(bool).sum()),
                "web_model_attempted_count": int(
                    group.web_model_attempted.astype(bool).sum()
                ),
                "safety_consensus_valid_count": int(
                    group.safety_consensus_valid.astype(bool).sum()
                ),
                "full_analysis_complete_count": int(
                    group.full_analysis_complete.astype(bool).sum()
                ),
                **{f"n_{state}": states[state] for state in LABELS},
                "coverage_status": "candidates_with_remaining_evidence_gates"
                if len(group)
                else "no_candidate_identified_in_this_release",
                "remaining_work": "Assay qualification and material/exposure review remain. No wet-lab validation was performed."
                if len(group)
                else "Primary-literature/pathway investigation and assay qualification required; no supported compound proposed.",
            }
        )
    return pd.DataFrame(rows)


def assay_plan(intents: list[dict], matrix: pd.DataFrame) -> pd.DataFrame:
    special = {
        "TRPV1": (
            "Human TRPV1 calcium response; confirm channel current",
            "Agonist challenge, vehicle, validated antagonist, parental cells",
        ),
        "AR": (
            "Human AR transactivation under agonist challenge; orthogonal target-gene response",
            "DHT agonist challenge, vehicle, authenticated antagonist, receptor-negative cells",
        ),
        "TYR": (
            "Human TYR concentration-response; orthogonal product readout; human melanocyte melanin",
            "Active-enzyme, enzyme-free compound blanks, substrate control, authenticated human-TYR inhibitor",
        ),
        "MMP1": (
            "Active human MMP-1 enzymatic assay and collagen degradation",
            "GM6001 reference, vehicle, no-enzyme and compound-only blanks; preserve Zn/Ca conditions",
        ),
        "MMP3": (
            "Active human MMP-3 enzyme assay and orthogonal substrate cleavage",
            "Qualified MMP inhibitor, vehicle, no-enzyme and compound-only blanks; preserve Zn/Ca conditions",
        ),
        "KLK5": (
            "Human KLK5 enzyme concentration-response; orthogonal substrate and CAMP cleavage",
            "Vehicle, enzyme-free blank, authenticated active and weak-binding controls; weak is not inactive",
        ),
        "NFE2L2": (
            "ARE reporter plus NRF2 dependence, downstream genes and oxidative-damage endpoint",
            "Vehicle, authenticated pathway inducer, NRF2 loss-of-function control, viability and reporter-interference controls",
        ),
        "COL1A1": (
            "Human dermal fibroblast procollagen-I propeptide and deposited collagen",
            "Vehicle, qualified positive control, viable-cell normalization; monitor fibrotic response",
        ),
        "FLG": (
            "Reconstructed epidermis: profilaggrin/processed filaggrin and localization",
            "Matched vehicle, tissue-integrity control; distinguish precursor and mature products",
        ),
        "CLDN1": (
            "Junctional claudin-1 imaging plus TEER/permeability",
            "Matched vehicle, integrity control; abundance alone is insufficient",
        ),
        "CAMP": (
            "Mature LL-37 separately from precursor CAMP and cleavage products",
            "Authenticated mature-peptide standard, precursor control, viable-cell normalization",
        ),
    }
    endpoint = {
        "TI-014-HA": (
            "Hyaluronan amount and molecular-size distribution",
            "Vehicle, assay standard and recovery controls",
        ),
        "TI-019-8OHDG": (
            "LC-MS/MS 8-OHdG with internal standard and DNA-normalized denominator",
            "Oxidative challenge, blank extraction, recovery and viability controls",
        ),
        "TI-025-CERAMIDE": (
            "LC-MS ceramide subclasses/chain lengths and barrier model response",
            "Vehicle, internal standards and viability controls",
        ),
        "TI-027-TEWL": (
            "Barrier-model integrity first; standardized TEWL after independent human-study approval",
            "Composition-matched vehicle, baseline, randomized sites and environmental equilibration",
        ),
    }
    rows = []
    for intent in intents:
        readout, controls = special.get(
            intent["gene_symbol"],
            endpoint.get(
                intent["intent_id"],
                (
                    intent["readout"]
                    + "; orthogonal assay in the intended skin compartment",
                    "Vehicle, assay-qualified positive/negative controls and viability control",
                ),
            ),
        )
        candidates = matrix[
            matrix.intent_id.eq(intent["intent_id"]) & matrix.is_curated_working_set
        ]
        rows.append(
            {
                "intent_id": intent["intent_id"],
                "biomarker": intent["biomarker"],
                "desired_effect": intent["desired_effect"],
                "model_context": intent["skin_compartment"],
                "proposed_readout": readout,
                "controls": controls,
                "candidate_names_and_states": " | ".join(
                    f"{r.compound_name}: {r.decision}" for r in candidates.itertuples()
                ),
                "stage": "proposed; not executed",
                "start_gate": "Resolve candidate exclusion/safety flags, authenticate lot/isomer/purity and qualify assay before testing. No human administration is authorized by this package.",
                "design": "Pilot concentration-response within measured solubility and non-cytotoxic range; randomize and blind wells/samples. Use independent biological repeats; estimate variance before powering confirmation.",
                "interpretation": "Prespecify the functional endpoint; record raw data, exclusions and controls. Docking and a computed PASS do not establish efficacy or safety.",
                "data_fields": "intent_id,compound_id,lot,purity,concentration,unit,vehicle,model,donor_or_batch,run_id,replicate_type,time,raw_signal,viability,control,QC_status,exclusion_reason",
            }
        )
    return pd.DataFrame(rows)


def spreadsheet_safe(frame: pd.DataFrame) -> pd.DataFrame:
    # Source assay text is untrusted spreadsheet input. Preserve CSV values but
    # prevent formulas when the working workbook is opened by the experiment team.
    def cell(value: object) -> object:
        if isinstance(value, (dict, list)):
            value = js(value)
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    return frame.map(cell)


def write_workbook(path: Path, tables: dict[str, pd.DataFrame]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in tables.items():
            spreadsheet_safe(frame).to_excel(writer, index=False, sheet_name=name)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="183D44")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[1].height = 32
            for column in range(1, len(frame.columns) + 1):
                name_length = len(str(frame.columns[column - 1]))
                sheet.column_dimensions[get_column_letter(column)].width = min(
                    48, max(18, name_length + 3)
                )
                if str(frame.columns[column - 1]) in {
                    "source_url",
                    "identity_source_url",
                    "fulltext_source_url",
                }:
                    for cells in sheet.iter_rows(
                        min_row=2, min_col=column, max_col=column
                    ):
                        cell = cells[0]
                        if isinstance(cell.value, str) and cell.value.startswith(
                            "https://"
                        ):
                            cell.hyperlink = cell.value
                            cell.style = "Hyperlink"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--evidence-dir", type=Path, required=True)
    parser.add_argument("--safety", type=Path, nargs="+", required=True)
    parser.add_argument("--literature-dir", type=Path, required=True)
    parser.add_argument("--review-attestation", type=Path, required=True)
    parser.add_argument("--verification-files", type=Path, nargs="*", default=[])
    parser.add_argument(
        "--curation",
        type=Path,
        default=ROOT / "data/curation/target_candidate_literature_20260915.json",
    )
    parser.add_argument(
        "--intents", type=Path, default=_private_input("SKINSCOUT_BIOMARKER_XLSX")
    )
    parser.add_argument(
        "--gene-map", type=Path, default=_private_input("SKINSCOUT_GENE_MAP_PATH")
    )
    parser.add_argument(
        "--historical",
        type=Path,
        default=_private_input("SKINSCOUT_DISCOVERY_CANDIDATES"),
    )
    parser.add_argument("--structural-dir", type=Path)
    parser.add_argument(
        "--safety-archive-dir",
        type=Path,
        help="Optional complete safety artifact tree, including pre-reconciliation batches",
    )
    parser.add_argument("--out-dir", type=Path, required=True)
    args = parser.parse_args()
    out = args.out_dir
    if out.exists() and any(out.iterdir()):
        raise FileExistsError(
            "Use a new empty output directory to preserve prior sealed releases"
        )
    out.mkdir(parents=True, exist_ok=True)
    if args.intents is None or args.gene_map is None:
        raise SystemExit(
            "--intents/--gene-map 또는 SKINSCOUT_BIOMARKER_XLSX/SKINSCOUT_GENE_MAP_PATH 환경변수가 필요합니다"
        )
    intents = build_target_intents(args.intents, args.gene_map)
    sources = pd.read_parquet(args.evidence_dir / "candidate_source_summary.parquet")
    evidence_manifest = json.loads((args.evidence_dir / "manifest.json").read_text())
    for name in ("candidate_source_summary.parquet", "candidate_evidence.parquet"):
        expected = evidence_manifest["artifacts"][name]["sha256"]
        if digest(args.evidence_dir / name) != expected:
            raise ValueError(f"Source evidence artifact hash mismatch: {name}")
    safety, safety_artifacts = load_safety(args.safety)
    curated = load_curated(
        args.curation, args.literature_dir, intents, args.review_attestation
    )
    control_records, control_frame, structural_artifacts = load_structural_controls(
        args.structural_dir
    )
    control_ranking = control_ranking_summary(control_frame)
    historical = pd.read_csv(args.historical)
    matrix, compounds = build_matrix(
        intents, sources, historical, curated, safety, control_records
    )
    lineage_counts = historical_counts(intents, historical, matrix)
    for control in control_records:
        mask = matrix.intent_id.eq(control["intent_id"]) & matrix.compound_id.eq(
            control["compound_id"]
        )
        matrix.loc[mask, "structure_status"] = control["structure_status"]
    coverage = coverage_table(intents, matrix)
    assays = assay_plan(intents, matrix)
    selected = select_priority(matrix)
    working = (
        matrix[matrix.is_curated_working_set]
        .copy()
        .sort_values(["intent_id", "compound_name"])
    )
    for name, frame in {
        "target_intents": pd.DataFrame(intents),
        "candidate_summary": compounds,
        "target_candidate_matrix": matrix,
        "analysis_coverage": coverage,
        "assay_plan": assays,
        "working_candidates": working,
        "experiment_priority": selected,
        "curated_evidence": pd.DataFrame(curated),
        "structural_controls": control_frame,
    }.items():
        serial = frame.copy()
        for column in serial:
            serial[column] = serial[column].map(
                lambda v: js(v) if isinstance(v, (dict, list)) else v
            )
        serial.to_csv(out / f"{name}.csv", index=False, encoding="utf-8-sig")
    matrix.to_parquet(out / "target_candidate_matrix.parquet", index=False)
    readme_rows = pd.DataFrame(
        [
            {
                "item": "Purpose",
                "value": "SkinScout first experiment-team handoff; proposed evidence validation, not experimentally confirmed compounds",
            },
            {
                "item": "Working subset",
                "value": f"{len(working)} curated target-compound pairs; full {len(matrix)}-pair matrix is in target_candidate_matrix.csv/parquet",
            },
            {
                "item": "Decision",
                "value": "experiment_priority / needs_evidence / exclude / unsupported; no automatic upgrade for FLAG_HIGH or missing data",
            },
            {
                "item": "PASS",
                "value": "Computational skin-sensitization screening only; does not establish experimental safety",
            },
            {
                "item": "Human use",
                "value": "No human administration, purchase, regulatory clearance or experimental execution was performed",
            },
        ]
    )
    safety_frame = compounds[compounds.compound_id.isin(safety)].copy()
    working_columns = [
        "category",
        "biomarker",
        "compound_name",
        "role",
        "decision_label",
        "safety_decision",
        "safety_run_valid",
        "safety_domain",
        "safety_missing_models",
        "direction_relation",
        "reviewed_evidence",
        "reviewed_conditions",
        "source_urls",
        "evidence_limitations",
        "next_experiment",
        "decision_reasons",
        "applicability",
        "analysis_scope_supported",
        "structure_status",
        "procurement_status",
        "formulation_exposure_status",
        "compound_id",
        "computed_inchikey",
        "canonical_isomeric_smiles",
    ]
    write_workbook(
        out / "SkinScout_experiment_handoff.xlsx",
        {
            "READ_ME": readme_rows,
            "Working_candidates": working[working_columns],
            "Priority": selected,
            "Coverage_33": coverage,
            "Safety_reanalysis": safety_frame,
            "Assay_plan": assays,
            "Primary_evidence": pd.DataFrame(curated),
            "Structural_controls": control_frame,
        },
    )
    with Chem.SDWriter(str(out / "working_candidate_structures.sdf")) as writer:
        for row in working.drop_duplicates("compound_id").itertuples():
            molecule = Chem.MolFromSmiles(row.canonical_isomeric_smiles)
            molecule.SetProp("_Name", row.compound_name)
            molecule.SetProp("compound_id", row.compound_id)
            molecule.SetProp("computed_inchikey", row.computed_inchikey)
            molecule.SetProp(
                "coordinate_status", "identity record; no experimental pose"
            )
            writer.write(molecule)
    evidence_out = out / "evidence"
    evidence_out.mkdir(exist_ok=True)
    for name in (
        "candidate_evidence.parquet",
        "manifest.json",
        "evidence_coverage.csv",
    ):
        shutil.copy2(args.evidence_dir / name, evidence_out / name)
    portable_paths = []
    if args.safety_archive_dir:
        copy_artifact_tree(args.safety_archive_dir, out / "safety", out, portable_paths)
    else:
        for index, path in enumerate(args.safety):
            safety_out = out / "safety" / f"batch_{index + 1}"
            copy_artifact_tree(path.parent, safety_out, out, portable_paths)
    if args.structural_dir:
        copy_artifact_tree(args.structural_dir, out / "structural", out, portable_paths)
    shutil.copy2(args.curation, out / "literature_curation.json")
    shutil.copy2(args.review_attestation, out / "literature_independent_review.json")
    review_notes = args.review_attestation.with_suffix(".md")
    if review_notes.exists():
        shutil.copy2(review_notes, out / "LITERATURE_REVIEW.md")
    if len({p.name for p in args.verification_files}) != len(args.verification_files):
        raise ValueError("Verification files require distinct names")
    for path in args.verification_files:
        (out / "verification").mkdir(exist_ok=True)
        shutil.copy2(path, out / "verification" / path.name)
    copy_artifact_tree(
        args.literature_dir / "identity", out / "identity", out, portable_paths
    )
    write_json(
        out / "portable_path_map.json",
        {
            "interpretation": "Original manifests retain provenance paths. This map resolves them to portable package filenames; file bytes/hashes are unchanged.",
            "files": portable_paths,
        },
    )
    source_manifest = json.loads(
        (args.literature_dir / "sources/source_manifest.json").read_text()
    )
    write_json(out / "primary_source_manifest.json", source_manifest)
    safety_metrics = safety_coverage_metrics(matrix)
    counts = {
        **lineage_counts,
        **safety_metrics,
        "original_rows": len({r["source_row_id"] for r in intents}),
        "expanded_intents": len(intents),
        "protein_ids": len({r["uniprot_id"] for r in intents if r["uniprot_id"]}),
        "protein_ids_with_local_evidence": int(
            (coverage.source_evidence_rows.gt(0) & coverage.uniprot_id.notna()).sum()
        ),
        "source_evidence_rows": int(sources.evidence_count.sum()),
        "target_compound_pairs": len(matrix),
        "unique_compounds": len(compounds),
        "unresolved_structure_ids": int(
            compounds.canonical_isomeric_smiles.isna().sum()
        ),
        "source_identity_mismatch_evidence_rows": int(
            matrix.source_identity_mismatch_count.sum()
        ),
        "compounds_without_safety_record": int(
            (~compounds.safety_record.astype(bool)).sum()
        ),
        "compounds_without_web_model_run": int(
            (~compounds.web_model_attempted.astype(bool)).sum()
        ),
        "curated_pairs": len(working),
        "curated_compounds": working.compound_id.nunique(),
        "experiment_priority_pairs": int(
            matrix.decision.eq("experiment_priority").sum()
        ),
        "displayed_priority_pairs": len(selected),
        "decisions": dict(Counter(matrix.decision)),
        "current_code_safety_decisions": dict(
            Counter(r["decision"] for r in safety.values())
        ),
        "structural_control_pairs": len(control_records),
        "structural_control_states": dict(
            Counter(r["structure_status"] for r in control_records)
        ),
        "safety_model_execution": dict(
            Counter(
                f"{name}:{result.get('execution', 'unknown')}"
                for record in safety.values()
                for name, result in record.get("model_results", {}).items()
            )
        ),
    }
    columns = [
        "biomarker",
        "compound_name",
        "role",
        "decision_label",
        "safety_decision",
        "direction_relation",
    ]
    table = working[columns].to_markdown(index=False)
    safety_record_metric = counts["safety_record_count"]
    web_model_metric = counts["web_model_attempted_count"]
    consensus_metric = counts["safety_consensus_valid_count"]
    complete_metric = counts["full_analysis_complete_count"]
    report = (
        "# SkinScout 실험팀 인계 — 2026-09-15\n\n"
        "이 자료는 계산·문헌 근거를 바탕으로 한 1차 실험 검토 자료입니다. 실험으로 효능·안전성이 확정된 후보는 없습니다.\n\n"
        f"원본 **{counts['original_rows']}개 항목 / {counts['expanded_intents']}개 목적 / {counts['protein_ids']}개 단백질 ID**를 보존했습니다. "
        f"원천 근거 {counts['source_evidence_rows']:,}행과 {len(matrix):,}개 표적–화합물 관계를 수집했습니다. "
        f"안전성 record는 {safety_record_metric['pairs']}쌍/{safety_record_metric['compounds']}종, "
        f"웹 모델 실행 시도는 {web_model_metric['compounds']}종, "
        f"합의 유효(부분 2-positive HALT 포함)는 {consensus_metric['pairs']}쌍/{consensus_metric['compounds']}종, "
        f"완전 분석은 {complete_metric['pairs']}쌍/{complete_metric['compounds']}종입니다. "
        f"record가 없는 고유 화합물은 {counts['compounds_without_safety_record']:,}종, "
        f"웹 모델을 실행하지 않은 고유 화합물은 {counts['compounds_without_web_model_run']:,}종입니다. "
        "모든 문헌을 망라한 탐색이나 전체 후보의 실험 검증을 뜻하지 않습니다.\n\n"
        "안전성은 최신 코드로 재평가했습니다. 실시간 조회와 정확 구조에 연결된 저장 응답 재해석을 구분하며, "
        "STopTox 장애로 확보하지 못한 신규 응답은 결측으로 남겼습니다. 각 방식·원래 시각·해시는 safety 파일에 기록했습니다.\n\n"
        f"현재 엄격한 실험 우선 조건을 충족한 관계는 **{counts['experiment_priority_pairs']}개**입니다. "
        "아래 작업 목록은 추가 검증할 소재와 대조군입니다. 제외·지원 밖 상태는 표기된 제한을 해소하기 전 실험 후보로 승격하지 않습니다.\n\n"
        + table
        + "\n\n"
        "## 파일 사용 순서\n\n"
        "1. `SkinScout_experiment_handoff.xlsx`: 작업 후보, 전체 33개 목적의 상태, 안전성 및 실험안.\n"
        "2. `target_candidate_matrix.csv/parquet`: 작업 목록 밖의 후보까지 포함한 전체 관계와 판정 사유.\n"
        "3. `evidence/candidate_evidence.parquet`, `curated_evidence.csv`: 개별 출처·시험·부등호·종·정확한 구조.\n"
        "4. `safety/`, `structural/`: 실제 실행·실패·보류와 근거 파일.\n"
        "5. `decision_manifest.json`, `checksums.sha256`: 개수, 입력·코드·출력 해시.\n\n"
        "## 실험팀이 먼저 확인할 사항\n\n"
        "실험안은 제안이며 수행되지 않았습니다. 후보의 현 상태와 근거를 검토하고, 정확한 이성질체·공급 lot·순도, 용해도, "
        "세포독성, 측정 간섭과 제형·노출을 확인해야 합니다. 보고된 농도는 각 논문의 조건으로, 사람 사용량을 뜻하지 않습니다. "
        "실험용 대조군과 소재 후보는 역할을 구분했습니다. CosIng 등재를 허가 또는 안전성 검증으로 취급하지 않았습니다.\n\n"
        "도킹은 지정 표적의 구조 실행 진단입니다. TYR 구리·MMP 아연 지원이 검증되지 않아 해당 도킹은 보류됩니다. "
        "HA·8-OHdG·Ceramide·TEWL은 endpoint이며 단백질 도킹을 하지 않습니다. "
        "AR의 DHT는 억제 목적과 반대인 작용제 대조군입니다. NRF2 경로 신호와 실제 항산화 효능·안전성은 구분해야 합니다.\n"
    )
    if control_ranking["status"] == "single_control_pair_only":
        mismatches = [
            metric
            for metric, match in control_ranking[
                "metric_order_matches_measured_affinity"
            ].items()
            if match is False
        ]
        report += "\nKLK5 대조군 두 개의 실제 도킹을 완료했습니다. "
        if mismatches:
            report += (
                f"실측보다 약한 대조군을 높게 평가한 지표: {', '.join(mismatches)}. "
            )
        report += "이 대조쌍 하나로 활성 순위 정확성을 검증하지 않으며, 도킹 점수로 소재 후보를 승격하지 않습니다. 원시 점수는 `structural_controls.csv`에 있습니다.\n"
    (out / "START_HERE.md").write_text(report, encoding="utf-8")
    reproduce = (
        "# Reproduce this release\n\n"
        "The full SkinScout checkout, its existing environment and source datasets are required. "
        "No new package installation is required. Use a new output directory to preserve this sealed release.\n\n"
        "The ordered commands and input paths are recorded in decision_manifest.json. "
        "First build target evidence, then re-evaluate the bounded safety queue, then execute the explicit structural pairs, "
        "and finally run scripts/build_target_lab_package.py with these artifacts. "
        "Source database lookup is complete for the configured targets; safety modeling is deliberately a bounded batch.\n\n"
        "Verify portable files with `sha256sum -c checksums.sha256` inside the extracted handoff directory. "
        "The source database and bibliography manifests retain original retrieval paths as provenance; "
        "portable_path_map.json maps original safety/structure/identity artifact paths to Windows-compatible filenames in this package; file contents and hashes are preserved. "
        "primary-paper URLs are in curated_evidence.csv. Original database/abstract retrieval files are retained in the SkinScout workspace.\n"
    )
    (out / "REPRODUCE.md").write_text(reproduce, encoding="utf-8")
    body = "<h1>SkinScout 실험팀 인계</h1><p>2026-09-15 · 계산·문헌 기반 1차 검토 자료. 실험 검증 완료 후보 없음.</p>"
    body += f"<p>원본 {counts['original_rows']}항목 / {counts['expanded_intents']}목적 · 엄격한 실험 우선 {counts['experiment_priority_pairs']}개</p>"
    body += '<p><a href="SkinScout_experiment_handoff.xlsx">엑셀 열기</a> · <a href="START_HERE.md">설명</a> · <a href="working_candidates.csv">작업 후보 CSV</a></p>'
    body += working[columns].to_html(index=False, escape=True)
    body += "<h2>전체 목적의 상태</h2>" + coverage[
        [
            "biomarker",
            "intent_id",
            "candidate_count",
            "curated_pair_count",
            "safety_record_count",
            "web_model_attempted_count",
            "safety_consensus_valid_count",
            "full_analysis_complete_count",
            "n_experiment_priority",
        ]
    ].to_html(index=False, escape=True)
    body += "<p>추가 근거 필요·제외·지원 밖 항목은 자동 승격하지 않습니다. 논문 농도는 사람 사용량이 아닙니다. 자세한 근거와 제한은 엑셀/CSV에 있습니다.</p>"
    (out / "START_HERE.html").write_text(
        '<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
        "<title>SkinScout 실험팀 인계</title><style>body{font-family:system-ui,sans-serif;max-width:1400px;margin:40px auto;padding:0 24px;color:#193c43}table{border-collapse:collapse;width:100%;font-size:14px}th,td{border:1px solid #ccd8dc;padding:9px;text-align:left}th{background:#e5eff0}a{color:#156c91}</style>"
        + body
        + "</html>",
        encoding="utf-8",
    )
    code = [
        Path(__file__),
        ROOT / "scripts/target_intent.py",
        ROOT / "scripts/compound_applicability.py",
        ROOT / "scripts/build_target_candidate_evidence.py",
        ROOT / "scripts/reanalyze_target_candidate_safety.py",
        ROOT / "scripts/screen_target_candidates.py",
    ]
    manifest = {
        "schema_version": SCHEMA,
        "created_at": datetime.now(UTC).isoformat(),
        "counts": counts,
        "structural_control_ordering": control_ranking,
        "scope": "First handoff: complete local-source lookup and explicit status coverage; bounded primary curation and current-code safety reanalysis with per-model live/historical provenance; no wet-lab, purchase or human administration.",
        "command_argv": sys.argv,
        "runtime": {"python": sys.version, "pandas": pd.__version__},
        "inputs": [
            {"path": str(p), "sha256": digest(p)}
            for p in [
                args.intents,
                args.gene_map,
                args.historical,
                args.curation,
                args.review_attestation,
                args.evidence_dir / "manifest.json",
                args.evidence_dir / "candidate_source_summary.parquet",
            ]
            + (
                [
                    args.structural_dir / "structural_manifest.json",
                    args.structural_dir / "pair_status.csv",
                    args.structural_dir / "source_record_manifest.json",
                ]
                if args.structural_dir
                else []
            )
        ],
        "code": [{"path": str(p), "sha256": digest(p)} for p in code],
        "safety_batches": [
            {"path": r["path"], "sha256": r["sha256"]} for r in safety_artifacts
        ],
        "structural_screen": structural_artifacts,
        "verification_evidence": [
            {"path": str(p), "sha256": digest(p)} for p in args.verification_files
        ],
        "policy": "No unreviewed source summary establishes minimum functional evidence. Unknown material eligibility, incomplete conditions, FLAG_HIGH or missing safety prohibit experiment_priority. No cross-assay potency score.",
        "outputs": [
            {
                "path": str(p.relative_to(out)),
                "bytes": p.stat().st_size,
                "sha256": digest(p),
            }
            for p in sorted(out.rglob("*"))
            if p.is_file()
            and p.name not in {"decision_manifest.json", "checksums.sha256"}
        ],
    }
    write_json(out / "decision_manifest.json", manifest)
    paths = [
        p
        for p in sorted(out.rglob("*"))
        if p.is_file() and p.name != "checksums.sha256"
    ]
    (out / "checksums.sha256").write_text(
        "".join(f"{digest(p)}  {p.relative_to(out)}\n" for p in paths), encoding="utf-8"
    )
    archive = out.parent / "SkinScout_experiment_handoff_20260915.zip"
    with zipfile.ZipFile(
        archive, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6
    ) as zipped:
        for path in sorted(out.rglob("*")):
            if path.is_file():
                zipped.write(path, arcname=str(Path(out.name) / path.relative_to(out)))
    print(
        js(
            {
                "counts": counts,
                "handoff": str(out),
                "archive": str(archive),
                "archive_sha256": digest(archive),
            }
        )
    )


if __name__ == "__main__":
    main()
